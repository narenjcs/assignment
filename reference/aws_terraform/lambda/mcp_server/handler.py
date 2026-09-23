"""
AWS MCP Server — Lambda-based Model Context Protocol Server
Exposes tools to Bedrock Agent and Orchestrator:
  - extract_docx_text   : Extract text from DOCX in S3
  - store_summary       : Persist summary to S3 + DynamoDB
  - get_document_metadata: Retrieve document info
  - invoke_bedrock_model: Direct model invocation for enrichment
"""
import json
import boto3
import os
import re
from datetime import datetime, timezone

s3_client      = boto3.client("s3")
bedrock_client = boto3.client("bedrock-runtime")
dynamodb       = boto3.resource("dynamodb")

DOC_BUCKET = os.environ["DOC_BUCKET"]
JOBS_TABLE = os.environ["JOBS_TABLE"]
MODEL_ID   = os.environ["MODEL_ID"]

# ── MCP Tool Registry ─────────────────────────────────────
MCP_TOOLS = {
    "extract_docx_text":      "Extract and return full text content from a DOCX file in S3",
    "store_summary":          "Store document summary to S3 results prefix and update DynamoDB",
    "get_document_metadata":  "Get metadata about a document from DynamoDB job record",
    "invoke_bedrock_model":   "Invoke Claude model directly for text enrichment tasks",
    "list_available_tools":   "List all available MCP tools and their descriptions",
}

def extract_docx_text(s3_bucket, s3_key):
    """Extract text from DOCX using python-docx (via Lambda layer)"""
    try:
        import docx  # python-docx from Lambda layer
        obj = s3_client.get_object(Bucket=s3_bucket, Key=s3_key)
        content = obj["Body"].read()
        
        # Write to /tmp for python-docx
        tmp_path = f"/tmp/{s3_key.split('/')[-1]}"
        with open(tmp_path, "wb") as f:
            f.write(content)
        
        doc  = docx.Document(tmp_path)
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        
        # Extract tables
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                if row_text.strip():
                    tables_text.append(row_text)
        
        if tables_text:
            text += "\n\n[TABLES]\n" + "\n".join(tables_text)
        
        return {
            "success":    True,
            "text":       text,
            "char_count": len(text),
            "s3_key":     s3_key
        }
    except Exception as e:
        return {"success": False, "error": str(e), "s3_key": s3_key}

def store_summary(job_id, summary, doc_type, s3_key, metadata=None):
    """Store summary to S3 and update DynamoDB"""
    result_key = f"results/{job_id}/summary.json"
    result_data = {
        "jobId":     job_id,
        "docType":   doc_type,
        "s3Key":     s3_key,
        "summary":   summary,
        "metadata":  metadata or {},
        "storedAt":  datetime.now(timezone.utc).isoformat(),
        "processor": "AWS_BEDROCK_AGENT"
    }
    
    s3_client.put_object(
        Bucket=DOC_BUCKET,
        Key=result_key,
        Body=json.dumps(result_data, indent=2),
        ContentType="application/json"
    )
    
    table = dynamodb.Table(JOBS_TABLE)
    table.update_item(
        Key={"jobId": job_id},
        UpdateExpression="SET resultS3Key = :k, summaryPreview = :p, updatedAt = :t",
        ExpressionAttributeValues={
            ":k": result_key,
            ":p": summary[:500],
            ":t": datetime.now(timezone.utc).isoformat()
        }
    )
    
    return {"success": True, "result_key": result_key, "job_id": job_id}

def get_document_metadata(job_id):
    """Retrieve job metadata from DynamoDB"""
    table = dynamodb.Table(JOBS_TABLE)
    resp  = table.get_item(Key={"jobId": job_id})
    item  = resp.get("Item", {})
    return {"success": True, "metadata": item} if item else {"success": False, "error": "Job not found"}

def invoke_bedrock_model(prompt, system_prompt=None, max_tokens=4096):
    """Direct Bedrock model invocation for enrichment"""
    messages = [{"role": "user", "content": prompt}]
    body = {
        "anthropic_version": "bedrock-2023-05-31",
        "max_tokens":        max_tokens,
        "messages":          messages
    }
    if system_prompt:
        body["system"] = system_prompt
    
    response = bedrock_client.invoke_model(
        modelId=MODEL_ID,
        body=json.dumps(body),
        contentType="application/json",
        accept="application/json"
    )
    result = json.loads(response["body"].read())
    return {
        "success": True,
        "text":    result["content"][0]["text"],
        "usage":   result.get("usage", {})
    }

# ── MCP Protocol Handler ──────────────────────────────────
def lambda_handler(event, context):
    """
    MCP Server entry point.
    Supports both direct tool calls and Bedrock Agent action group format.
    """
    print(f"[MCP-SERVER] Event: {json.dumps(event)[:500]}")
    
    # Bedrock Agent Action Group format
    if "actionGroup" in event:
        tool_name  = event.get("function", "")
        parameters = {p["name"]: p["value"] for p in event.get("parameters", [])}
        result     = dispatch_tool(tool_name, parameters)
        return {
            "actionGroup":        event["actionGroup"],
            "function":           tool_name,
            "functionResponse":   {
                "responseBody": {
                    "TEXT": {"body": json.dumps(result)}
                }
            }
        }
    
    # Direct Lambda invocation format
    tool_name  = event.get("tool", "")
    tool_input = event.get("input", {})
    result     = dispatch_tool(tool_name, tool_input)
    return {"statusCode": 200, "result": result}

def dispatch_tool(tool_name, params):
    if tool_name == "extract_docx_text":
        return extract_docx_text(
            params.get("s3_bucket", DOC_BUCKET),
            params["s3_key"]
        )
    elif tool_name == "store_summary":
        return store_summary(
            params["job_id"], params["summary"],
            params.get("doc_type", "DOCX"),
            params.get("s3_key", ""),
            params.get("metadata")
        )
    elif tool_name == "get_document_metadata":
        return get_document_metadata(params["job_id"])
    elif tool_name == "invoke_bedrock_model":
        return invoke_bedrock_model(
            params["prompt"],
            params.get("system_prompt"),
            params.get("max_tokens", 4096)
        )
    elif tool_name == "list_available_tools":
        return {"tools": MCP_TOOLS}
    else:
        return {"success": False, "error": f"Unknown tool: {tool_name}"}
